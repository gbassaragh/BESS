import json, openpyxl, numpy as np
from openpyxl.styles import Font, PatternFill, Alignment
B="/tmp/claude-0/-home-user-BESS/078f934e-4ca0-5169-be3f-cb66ff0e2514/scratchpad/build"
AJ=json.load(open(f"{B}/adjust_results.json")); NY=json.load(open(f"{B}/ny_cohort.json")); IXJ=json.load(open(f"{B}/indices.json")); P_=IXJ["indices"]
wb=openpyxl.load_workbook(f"{B}/Hyde_Park_BESS_Cost_Analysis.xlsx")
if "2026$ Adjustment" in wb.sheetnames: del wb["2026$ Adjustment"]
ws=wb.create_sheet("2026$ Adjustment", index=wb.sheetnames.index("Public Projects")+1)
NAVY="001E60"; hdr=Font(bold=True,color="FFFFFF",name="Calibri",size=10); hfill=PatternFill("solid",fgColor=NAVY)
base=Font(name="Calibri",size=10); blue=Font(name="Calibri",size=10,color="0000FF"); bold=Font(name="Calibri",size=10,bold=True); title=Font(name="Georgia",size=14,bold=True,color=NAVY); note=Font(name="Calibri",size=9,italic=True,color="6B7280")
def H(r,c,vals):
    for j,v in enumerate(vals):
        cell=ws.cell(row=r,column=c+j,value=v); cell.font=hdr; cell.fill=hfill; cell.alignment=Alignment(wrap_text=True,vertical="center")
ws["A1"]="Dollar-year adjustment of the public comparables to 2026 (deck slides 12 and 14-16; revised 5 Oct 2026)"; ws["A1"].font=title
ws["A2"]="Fixed quantities, not fixed shares: C(y) = s0 x I(y)/I(anchor) + (1 - s0) x HW(y)/HW(anchor); cost_2026 = cost_y x C(2026)/C(y). s0 = battery share at anchor-year prices (by owner type), I = battery index (selector B6), HW = Handy-Whitman North Atlantic. The battery share a project had in its own year is s0 x I(y)/I(anchor) / C(y) and is shown per row. Blue = inputs."; ws["A2"].font=note
ws["A4"]="Anchor battery share, utility-owned"; ws["B4"]=AJ["share_utility"]; ws["B4"].font=blue; ws["B4"].number_format="0%"
ws["A5"]="Anchor battery share, developer-owned"; ws["B5"]=AJ["share_developer"]; ws["B5"].font=blue; ws["B5"].number_format="0%"
ws["A6"]="Index selector (1 installed, 2 pack, 3 LCOS, 4 turnkey, 5 NREL)"; ws["B6"]=1; ws["B6"].font=blue
ws["A7"]="Anchor year for the shares"; ws["B7"]=AJ["anchor_year"]; ws["B7"].font=blue
r=9; ws.cell(row=r,column=1,value="A. Battery indices as published ($/kWh; LCOS $/MWh) and the construction index").font=bold; r+=1
H(r,1,["Year","Installed all-in (EIA 2015-22, LBNL 2024)","BNEF pack (real)","Lazard LCOS midpoint","BNEF turnkey","NREL base year","Handy-Whitman NA production plant","Note"]); r+=1
tod=lambda pts:{int(k):v for k,v in pts.items()}
inst=tod(P_["installed"]["points"]); pack=tod(P_["pack"]["points"]); lcos=tod(P_["lcos"]["points"]); turn=tod(P_["turnkey"]["points"]); nrel=tod(P_["nrel"]["points"]); hw=tod(IXJ["construction_index"]["points"])
def ip(d,y): xs=sorted(d); return float(np.interp(y,xs,[d[x] for x in xs]))
t0=r
for y in range(2015,2027):
    ws.cell(row=r,column=1,value=y).font=base
    for c,d in ((2,inst),(3,pack),(4,lcos),(5,turn),(6,nrel),(7,hw)):
        cell=ws.cell(row=r,column=c,value=round(ip(d,y),1)); cell.font=blue if y in d else base
    ws.cell(row=r,column=8,value=("published (blue)" if any(y in d for d in (inst,pack,lcos,turn,nrel,hw)) else "interpolated")).font=note; r+=1
t1=r-1
ws.cell(row=r,column=1,value=IXJ["note"]+" Installed: "+P_["installed"]["source"]+". Handy-Whitman: "+IXJ["construction_index"]["source"]).font=note
IDXCOL=lambda: f"CHOOSE($B$6,$B${t0}:$B${t1},$C${t0}:$C${t1},$D${t0}:$D${t1},$E${t0}:$E${t1},$F${t0}:$F${t1})"
r+=2; ws.cell(row=r,column=1,value=f"B. Reference class ({len(AJ['projects'])} comparables at or below 100 MWh, utility- and developer-owned)").font=bold; r+=1
H(r,1,["Project","COD year","Nominal $/kWh","Owner","Anchor share s0","Index at COD","Index at anchor","Index 2026","HW at COD","HW at anchor","HW 2026","Factor to 2026 (fixed quantities)","2026$ $/kWh","Battery share in its own year","Earlier fixed-share factor (40%)","Naive: whole cost at battery rate"]); r+=1
p0=r
for p in sorted(AJ["projects"],key=lambda x:x["nominal"]):
    ws.cell(row=r,column=1,value=p["project"]).font=base; ws.cell(row=r,column=2,value=p["cod_year"]).font=blue; c=ws.cell(row=r,column=3,value=round(p["nominal"],1)); c.font=blue; c.number_format='"$"#,##0'
    ws.cell(row=r,column=4,value=("utility" if p.get("utility")==1 else "developer")).font=base
    c=ws.cell(row=r,column=5,value=f'=IF(D{r}="utility",$B$4,$B$5)'); c.number_format="0%"
    ws.cell(row=r,column=6,value=f"=INDEX({IDXCOL()},MATCH(MIN(B{r},2026),$A${t0}:$A${t1},0))")
    ws.cell(row=r,column=7,value=f"=INDEX({IDXCOL()},MATCH($B$7,$A${t0}:$A${t1},0))")
    ws.cell(row=r,column=8,value=f"=INDEX({IDXCOL()},MATCH(2026,$A${t0}:$A${t1},0))")
    ws.cell(row=r,column=9,value=f"=INDEX($G${t0}:$G${t1},MATCH(MIN(B{r},2026),$A${t0}:$A${t1},0))")
    ws.cell(row=r,column=10,value=f"=INDEX($G${t0}:$G${t1},MATCH($B$7,$A${t0}:$A${t1},0))")
    ws.cell(row=r,column=11,value=f"=INDEX($G${t0}:$G${t1},MATCH(2026,$A${t0}:$A${t1},0))")
    f=ws.cell(row=r,column=12,value=f"=(E{r}*H{r}/G{r}+(1-E{r})*K{r}/J{r})/(E{r}*F{r}/G{r}+(1-E{r})*I{r}/J{r})"); f.number_format="0.00"
    a=ws.cell(row=r,column=13,value=f"=C{r}*L{r}"); a.number_format='"$"#,##0'
    sh=ws.cell(row=r,column=14,value=f"=E{r}*F{r}/G{r}/(E{r}*F{r}/G{r}+(1-E{r})*I{r}/J{r})"); sh.number_format="0%"
    fx=ws.cell(row=r,column=15,value=f"=0.4*H{r}/F{r}+0.6*K{r}/I{r}"); fx.number_format="0.00"
    nv=ws.cell(row=r,column=16,value=f"=C{r}*H{r}/F{r}"); nv.number_format='"$"#,##0'; r+=1
p1=r-1
ws.cell(row=r,column=1,value="Median").font=bold
for col in ("C","M","P"):
    m=ws[f"{col}{r}"]; m.value=f"=MEDIAN({col}{p0}:{col}{p1})"; m.font=bold; m.number_format='"$"#,##0'
m=ws[f"L{r}"]; m.value=f"=MEDIAN(L{p0}:L{p1})"; m.number_format="0.00"; r+=1
for lab,v in (("Hyde Park installed $833, percentile",833),("Hyde Park station complete direct $1,483, percentile",1483),("Hyde Park approval $2,862, percentile",2862)):
    ws.cell(row=r,column=1,value=lab).font=base
    for col in ("C","M","P"):
        c1=ws[f"{col}{r}"]; c1.value=f"=COUNTIF({col}{p0}:{col}{p1},\"<\"&{v})/COUNT({col}{p0}:{col}{p1})"; c1.number_format="0%"
    r+=1
ws.cell(row=r,column=1,value=f"Python model (adjust_model.py): installed index, owner-specific anchor shares: median ${AJ['indices']['installed']['median']:,.0f}, Hyde Park station percentile {AJ['indices']['installed']['hp_pct']['station_direct']:.0%}; earlier fixed-share form ${AJ['indices']['installed_fixed_share']['median']:,.0f}. Implied battery share of a utility-owned project: {AJ['indices']['installed']['share_2016']:.0%} in 2016, {AJ['indices']['installed']['share_2018']:.0%} in 2018.").font=note; r+=2
ws.cell(row=r,column=1,value="C. New York distribution-scale cohort (not in the regression; nominal dollars of the year shown)").font=bold; r+=1
H(r,1,["Project","Owner","MW","MWh","COD year","Cost $","$/kWh","$/kW","Hours","Cost basis","Confidence","Source"]); r+=1
for p in NY["projects"]:
    vals=[p["project"],p["owner"],p["mw"],p["mwh"],p["cod_year"],p["cost_usd"],f"=F{r}/D{r}/1000",f"=F{r}/C{r}/1000",f"=D{r}/C{r}",p["cost_basis"],p["confidence"],p["url"]]
    for j,v in enumerate(vals,1):
        c=ws.cell(row=r,column=j,value=v); c.font=blue if j in (3,4,5,6) else base
        if j in (6,7,8): c.number_format='"$"#,##0'
        if j==9: c.number_format="0.0"
    r+=1
for ix in NY["indices"]:
    ws.cell(row=r,column=1,value=ix["name"]).font=base; ws.cell(row=r,column=2,value="NYSERDA program average").font=base; ws.cell(row=r,column=5,value=ix["year"]).font=blue; c=ws.cell(row=r,column=7,value=ix["kwh"]); c.font=blue; c.number_format='"$"#,##0'; ws.cell(row=r,column=10,value=ix["source"]).font=base; ws.cell(row=r,column=11,value=ix["confidence"]).font=base; r+=1
ic=NY["interconnection"]; ws.cell(row=r,column=1,value="Con Edison two-part test (CESIR), NY-BEST / NYSEIA filing 11 Mar 2026, Case 24-E-0621").font=base; ws.cell(row=r,column=3,value=ic["mw"]).font=blue; ws.cell(row=r,column=6,value=ic["avg_upgrade_usd"]).font=blue; ws.cell(row=r,column=6).number_format='"$"#,##0'; ws.cell(row=r,column=10,value=ic.get("note","")).font=base; r+=1
ws.cell(row=r,column=1,value=NY["note"]).font=note
for col,wdt in zip("ABCDEFGHIJKLMNOP",(46,12,12,11,11,11,11,11,10,10,10,14,12,12,12,14)): ws.column_dimensions[col].width=wdt
ws.freeze_panes="A4"
rd=wb["README"]
rows=[i for i in range(1,rd.max_row+1) if rd.cell(row=i,column=1).value in ("2026$ Adjustment","Tab: 2026$ Adjustment")]
tgt=rows[-1] if rows else rd.max_row+1
rd.cell(row=tgt,column=1,value="Tab: 2026$ Adjustment").font=bold
rd.cell(row=tgt,column=2,value="Dollar-year adjustment of the reference-class comparables to 2026 under five battery indices (selector B6) with fixed quantities: anchor battery shares by owner (B4, B5, at the B7 anchor year) back-cast to each project's own year, Handy-Whitman on the remainder; medians, Hyde Park percentiles, the superseded fixed-share factor and the naive whole-cost deflation; New York cohort. Deck slides 12-16. Revised 5 Oct 2026.").font=base
wb.save(f"{B}/Hyde_Park_BESS_Cost_Analysis.xlsx"); print("saved", wb.sheetnames)
