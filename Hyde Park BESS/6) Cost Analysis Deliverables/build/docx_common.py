from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY=RGBColor(0x00,0x1E,0x60); GOLD=RGBColor(0xFF,0xC7,0x2C); BURG=RGBColor(0x6E,0x1F,0x2E); GRAY=RGBColor(0x59,0x59,0x59)
def new_doc():
    d=Document()
    s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11); s.left_margin=s.right_margin=Inches(1); s.top_margin=Inches(0.9); s.bottom_margin=Inches(0.9)
    st=d.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(11); st.element.rPr.rFonts.set(qn("w:eastAsia"),"Calibri")
    st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.08
    for lvl,size in ((1,16),(2,13),(3,11.5)):
        h=d.styles[f"Heading {lvl}"]; h.font.name="Georgia"; h.font.size=Pt(size); h.font.bold=True; h.font.color.rgb=NAVY
        h.element.rPr.rFonts.set(qn("w:eastAsia"),"Georgia"); h.element.rPr.rFonts.set(qn("w:ascii"),"Georgia"); h.element.rPr.rFonts.set(qn("w:hAnsi"),"Georgia")
        h.paragraph_format.space_before=Pt(14 if lvl==1 else 10); h.paragraph_format.space_after=Pt(4); h.paragraph_format.keep_with_next=True
    z=d.settings.element.find(qn('w:zoom'))
    if z is not None and z.get(qn('w:percent')) is None: z.set(qn('w:percent'),'100')
    return d
def shade(cell,hex_):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:color"),"auto"); shd.set(qn("w:fill"),hex_); tcPr.append(shd)
def set_cell_text(cell,text,bold=False,color=None,size=9.5,align=None,italic=False):
    cell.text=""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
    r=p.add_run(str(text)); r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    if color: r.font.color.rgb=color
    if align: p.alignment=align
def table(d,headers,rows,widths=None,num_cols=(),font=9.5,header_fill="001E60",zebra=True):
    t=d.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; set_cell_text(c,h,bold=True,color=RGBColor(0xFF,0xFF,0xFF),size=font); shade(c,header_fill)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            al=WD_ALIGN_PARAGRAPH.RIGHT if i in num_cols else None
            bold = isinstance(v,str) and v.startswith("**")
            txt = v[2:] if bold else v
            set_cell_text(cells[i],txt,size=font,align=al,bold=bold)
            if zebra and ri%2==1: shade(cells[i],"F2F4F8")
    if widths:
        t.autofit=False
        for i,w in enumerate(widths): t.columns[i].width=Inches(w)
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def para(d,text,bold=False,italic=False,size=None,color=None,after=6,style=None):
    p=d.add_paragraph(style=style) if style else d.add_paragraph()
    r=p.add_run(text); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=color
    p.paragraph_format.space_after=Pt(after); return p
def bullets(d,items,style="List Bullet"):
    for it in items:
        p=d.add_paragraph(style=style); p.paragraph_format.space_after=Pt(3)
        if isinstance(it,tuple):
            r=p.add_run(it[0]); r.bold=True; p.add_run(it[1])
        else: p.add_run(it)
def callout(d,label,text,fill="FFF3CC"):
    t=d.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.rows[0].cells[0]; shade(c,fill)
    c.text=""; p=c.paragraphs[0]; r=p.add_run(label+" "); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=NAVY
    r2=p.add_run(text); r2.font.size=Pt(10); p.paragraph_format.space_after=Pt(2)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
def footer(d,text):
    for s in d.sections:
        f=s.footer; p=f.paragraphs[0]; p.text=""; r=p.add_run(text); r.font.size=Pt(8); r.font.color.rgb=GRAY; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def m(v): return f"${v/1e6:,.2f}M"
def k(v): return f"${v/1e3:,.0f}K"
def d0(v): return f"${v:,.0f}"
