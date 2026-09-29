"""Heuristic text-fit QA for python-pptx decks: estimates wrapped line count per text frame with PIL font metrics
(Liberation Sans for Calibri/Arial - conservative, wider than Calibri; DejaVu Serif for Georgia - close) and flags frames
whose estimated text height exceeds the box height. Also flags shapes that run off the slide and overlapping text boxes."""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont
FONTS={"Calibri":"/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf","Georgia":"/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"}
BOLD={"Calibri":"/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf","Georgia":"/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"}
CAL_SCALE=0.92  # Calibri is ~8% narrower than Liberation Sans/Arial
def font(name,size,bold):
    path=(BOLD if bold else FONTS).get(name,FONTS["Calibri"]); return ImageFont.truetype(path,int(size*4))  # 4x for precision
def width(txt,name,size,bold):
    f=font(name,size,bold); w=f.getlength(txt)/4.0
    return w*(CAL_SCALE if name=="Calibri" else 1.0)
def est_height(tf,box_w_pt):
    total=0.0
    for p in tf.paragraphs:
        runs=[(r.text,r.font.name or "Calibri",(r.font.size.pt if r.font.size else 18),bool(r.font.bold)) for r in p.runs]
        if not runs: total+=14; continue
        size=max(r[2] for r in runs); ls=p.line_spacing if isinstance(p.line_spacing,float) else 1.0
        # word wrap greedy across runs
        words=[]; 
        for t,n,s,b in runs:
            for w in t.replace("\n"," ").split(" "):
                if w: words.append((w,n,s,b))
        lines=1; cur=0.0
        for w,n,s,b in words:
            ww=width(w+" ",n,s,b)
            if cur+ww>box_w_pt and cur>0: lines+=1; cur=ww
            else: cur+=ww
        sa=(p.space_after.pt if p.space_after else 0)
        total+=lines*size*1.2*ls+sa
    return total
prs=Presentation(sys.argv[1]); SW=prs.slide_width; SH=prs.slide_height
issues=0
for i,s in enumerate(prs.slides,1):
    boxes=[]
    for sh in s.shapes:
        if sh.left is None: continue
        if sh.left<0 or sh.top<0 or sh.left+sh.width>SW+Emu(9525) or sh.top+sh.height>SH+Emu(9525):
            print(f"slide {i}: OFF-SLIDE {sh.shape_type} '{(sh.text_frame.text[:30] if sh.has_text_frame else '')}'"); issues+=1
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf=sh.text_frame; bw=Emu(sh.width).pt-(tf.margin_left+tf.margin_right)/12700; bh=Emu(sh.height).pt-(tf.margin_top+tf.margin_bottom)/12700
            h=est_height(tf,bw)
            if h>bh*1.02:
                print(f"slide {i}: OVERFLOW est {h:.0f}pt > box {bh:.0f}pt  '{tf.text[:60].replace(chr(10),' / ')}'"); issues+=1
            boxes.append((sh.left,sh.top,sh.left+sh.width,sh.top+sh.height,tf.text[:25]))
        if sh.has_table:
            t=sh.table; est=0
            for r in t.rows:
                mx=0
                for c in r.cells:
                    cw=Emu(t.columns[list(r.cells).index(c)].width).pt-8
                    mx=max(mx,est_height(c.text_frame,cw)+6)
                est+=mx
            if est>Emu(sh.height).pt*1.05:
                print(f"slide {i}: TABLE taller than frame est {est:.0f}pt > {Emu(sh.height).pt:.0f}pt (rows {len(t.rows)}) -- check bottom margin"); issues+=1
            boxes.append((sh.left,sh.top,sh.left+sh.width,sh.top+Emu(int(max(sh.height,est*12700))),"TABLE"))
    for a in range(len(boxes)):
        for b in range(a+1,len(boxes)):
            A=boxes[a];Bx=boxes[b]
            ox=min(A[2],Bx[2])-max(A[0],Bx[0]); oy=min(A[3],Bx[3])-max(A[1],Bx[1])
            if ox>Emu(91440) and oy>Emu(91440):
                print(f"slide {i}: OVERLAP '{A[4]}' x '{Bx[4]}' ({Emu(ox).inches:.2f}in x {Emu(oy).inches:.2f}in)"); issues+=1
print("issues:",issues)
