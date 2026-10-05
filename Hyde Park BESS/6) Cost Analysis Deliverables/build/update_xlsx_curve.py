import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
B="/tmp/claude-0/-home-user-BESS/078f934e-4ca0-5169-be3f-cb66ff0e2514/scratchpad/build"
CV=json.load(open(f"{B}/curve.json")); ES=CV["escalation"]
wb=openpyxl.load_workbook(f"{B}/Hyde_Park_BESS_Cost_Analysis.xlsx")
if "Cost Curve" in wb.sheetnames: del wb["Cost Curve"]
ws=wb.create_sheet("Cost Curve", index=wb.sheetnames.index("Benchmarks")+1)
NAVY="001E60"; hdr=Font(bold=True,color="FFFFFF",name="Calibri",size=10); hfill=PatternFill("solid",fgColor=NAVY)
base=Font(name="Calibri",size=10); blue=Font(name="Calibri",size=10,color="0000FF"); bold=Font(name="Calibri",size=10,bold=True); title=Font(name="Georgia",size=14,bold=True,color=NAVY)
def H(r,c,vals):
    for j,v in enumerate(vals):
        cell=ws.cell(row=r,column=c+j,value=v); cell.font=hdr; cell.fill=hfill; cell.alignment=Alignment(wrap_text=True,vertical="center")
ws["A1"]="Battery installed-cost curve and escalation basis (deck slides 10-11)"; ws["A1"].font=title
ws["A2"]=CV["note"]; ws["A2"].font=Font(name="Calibri",size=9,italic=True,color="6B7280")
r=4; ws.cell(row=r,column=1,value="A. Industry series as published ($/kWh)").font=bold; r+=1
H(r,1,["Year","Series","$/kWh","Dollar basis","Scope","Confidence","Source"]); r+=1
first=r
for s in CV["series"]:
    for x,y in s["points"]:
        ws.cell(row=r,column=1,value=x).font=blue; ws.cell(row=r,column=2,value=s["name"]).font=base; ws.cell(row=r,column=3,value=y).font=blue
        ws.cell(row=r,column=4,value=s["basis"]).font=base; ws.cell(row=r,column=5,value=s["scope"]).font=base; ws.cell(row=r,column=6,value=s["confidence"]).font=base; ws.cell(row=r,column=7,value=s["url"]).font=base
        r+=1
last=r-1
r+=1; ws.cell(row=r,column=1,value="B. Hyde Park points ($/kWh at 20,000 kWh)").font=bold; r+=1
H(r,1,["Year (decimal)","Point","$/kWh","Basis"]); r+=1
hp=[(2023.5,"2023 indicative bids, low (Company 2 Opt 2 $9,586,600)","='2023 Indicative'!B3/20000" if False else 479,"July 2023 indicative sheet"),(2023.5,"2023 indicative bids, high (Company 3 $18,222,342)",911,"July 2023 indicative sheet"),(2026.2,"Nomad equipment only ($10,638,000)","=Layers!B25/20000*10638000/12106850","Cover agreement SOV"),(2026.2,"Nomad contract ($12,106,850)","=Layers!B25/20000","Cover agreement Exhibit C"),(2026.75,"Battery as carried in Rev 3, layer A (contract balance escalated + actuals)","=Layers!B5/20000","Rev 3 lines 432 + 445"),(2026.75,"Battery installed, layers A+B+C","=(Layers!B5+Layers!B6+Layers!B7)/20000","Rev 3 layers A, B, C")]
hp_first=r
for x,lab,v,bas in hp:
    ws.cell(row=r,column=1,value=x).font=blue; ws.cell(row=r,column=2,value=lab).font=base; c=ws.cell(row=r,column=3,value=v); c.font=base if isinstance(v,str) else blue; c.number_format='"$"#,##0'; ws.cell(row=r,column=4,value=bas).font=base; r+=1
r+=1; ws.cell(row=r,column=1,value="C. Adjusting a 4-h, 60-240 MWh index to 2-h, 20 MWh").font=bold; r+=1
H(r,1,["Input","Value","Note"]); r+=1
adj0=r
rows=[("Duration factor, low",CV["adjustment"]["duration_factor_lo"],"NREL cost structure: pack $/kWh x hours + BOS $/kW; 2-h vs 4-h"),("Duration factor, high",CV["adjustment"]["duration_factor_hi"],"as above"),("Size elasticity",CV["adjustment"]["size_elasticity"],"Appendix B regression, ln(MWh) coefficient"),("Reference MWh",CV["adjustment"]["ref_mwh"],"NREL 60 MW x 4 h"),("Hyde Park MWh",CV["adjustment"]["hp_mwh"],""),("Size factor",f"=(B{adj0+4}/B{adj0+3})^B{adj0+2}","(20/240)^-0.173"),("Combined factor, low",f"=B{adj0}*B{adj0+5}",""),("Combined factor, high",f"=B{adj0+1}*B{adj0+5}",""),("BNEF US turnkey 2025, adjusted low",f"=219*B{adj0+6}","vs Nomad equipment $532"),("BNEF US turnkey 2025, adjusted high",f"=219*B{adj0+7}",""),("NREL 2024 base $334 to 2028 at 3%/yr",f"=334*(1+B{adj0+13})^4",""),("NREL 2028, adjusted low",f"=B{adj0+10}*B{adj0+6}","vs Hyde Park installed $833"),("NREL 2028, adjusted high",f"=B{adj0+10}*B{adj0+7}",""),("Escalation rate",ES["rate"],"ESF template, approx.")]
for lab,v,note in rows:
    ws.cell(row=r,column=1,value=lab).font=base; c=ws.cell(row=r,column=2,value=v); c.font=blue if not (isinstance(v,str) and v.startswith("=")) else base; c.number_format='0.000' if (isinstance(v,float) or (isinstance(v,str) and 'factor' in lab.lower())) else '"$"#,##0' if 'adjusted' in lab or 'NREL 2024' in lab else '0.###'; ws.cell(row=r,column=3,value=note).font=base; r+=1
r+=1; ws.cell(row=r,column=1,value="D. Escalation in Rev 3 (ESF Overview row 34, included in directs)").font=bold; r+=1
H(r,1,["Item","$","$/kWh","Note"]); r+=1
e0=r
erows=[("Escalation, Station 360",ES["sta360"],"ESF Sta 360 Rev 3 Overview"),("Escalation, Station 496",ES["sta496"],"ESF Total Rev 3 Overview"),("Escalation, project 21334",f"=B{e0}+B{e0+1}","ESF Overview shows $2,948,100 after rounding"),("of which on the Nomad balance (fixed price)","=Layers!B26","Rev 3 line 432 less contract balance"),("of which on other scope",f"=B{e0+2}-B{e0+3}",""),("Share of approval level",f"=B{e0+2}/57247300",""),("Spend 2027-2028 (Schedule tab)",ES["spend"]["2027"]+ES["spend"]["2028"],"ESF Schedule tab, 2027 + 2028 columns"),("Escalation from a 12-month slip at the template rate",f"=B{e0+6}*B{adj0+13}","before AFUDC, PM, indirects"),("Non-battery material (Materials category less layer A)",f"=16678100-Layers!B5","exposure to tariff / commodity moves above 3%")]
for lab,v,note in erows:
    ws.cell(row=r,column=1,value=lab).font=base; c=ws.cell(row=r,column=2,value=v); c.font=blue if not (isinstance(v,str) and v.startswith("=")) else base; c.number_format='0.0%' if 'Share' in lab else '"$"#,##0'
    if 'Share' not in lab: d=ws.cell(row=r,column=3,value=f"=B{r}/20000"); d.font=base; d.number_format='"$"#,##0'
    ws.cell(row=r,column=4,value=note).font=base; r+=1
r+=1; ws.cell(row=r,column=1,value="E. Same comparison, same dollars ($/kWh)").font=bold; r+=1
H(r,1,["Boundary","Nominal as carried","Ex-escalation (2026$)","Benchmark brought to 2028 at the template rate","Note"]); r+=1
c0=r
crow=[("Vendor contract","=Layers!B25/20000","=B{r}",f"=219*(1+B{adj0+13})^3","BNEF US turnkey 2025 -> 2028"),("Battery as carried (A)","=Layers!B5/20000",f"=B{{r}}-C{e0+3}",f"=334*(1+B{adj0+13})^4","NREL 4-h 2024$ -> 2028"),("Battery installed (A+B+C)","=(Layers!B5+Layers!B6+Layers!B7)/20000",f"=B{{r}}-C{e0+3}-C{e0+4}*(Layers!B6+Layers!B7)/(Layers!B6+Layers!B7+Layers!B8+Layers!B9+Layers!B11+Layers!B12+Layers!B13)",f"=B{adj0+10}*B{adj0+6}","NREL 2028 adjusted to 2-h, 10 MW (low); high = adjusted high"),("Station complete, direct","='Unit Costs'!D10",f"=B{{r}}-C{e0}",f"=436*(1+B{adj0+13})^5","EIA/S&L 2023$ -> 2028")]
for lab,b,c,dv,note in crow:
    ws.cell(row=r,column=1,value=lab).font=base
    for col,v in ((2,b),(3,c.replace("{r}",str(r))),(4,dv)):
        cell=ws.cell(row=r,column=col,value=v); cell.font=base; cell.number_format='"$"#,##0'
    ws.cell(row=r,column=5,value=note).font=base; r+=1
ws.cell(row=r,column=1,value="Ex-escalation removes the Station 360 escalation from the direct boundaries; the B+C share is allocated pro rata to direct cost across Station 360 layers B to I and is approximate.").font=Font(name="Calibri",size=9,italic=True,color="6B7280")
for col,wdt in zip("ABCDEFG",(44,30,16,34,60,12,60)): ws.column_dimensions[col].width=wdt
ws.freeze_panes="A4"
# README sheet: append a line for the new tab
rd=wb["README"]
lastr=max(c.row for row in rd.iter_rows() for c in row if c.value is not None)
rd.cell(row=lastr+1,column=1,value="Cost Curve").font=bold
rd.cell(row=lastr+1,column=2,value="Industry installed-cost series 2015-2035 as published (EIA, BNEF, NREL), Hyde Park points, the 2-h / 10 MW adjustment, the escalation carried in Rev 3 and the common-dollar comparison (deck slides 10-11). Added 2 Oct 2026.").font=base
wb.save(f"{B}/Hyde_Park_BESS_Cost_Analysis.xlsx"); print("saved; sheets", wb.sheetnames)
