from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
NAVY=RGBColor(0x00,0x1E,0x60); GOLD=RGBColor(0xFF,0xC7,0x2C); BURG=RGBColor(0x6E,0x1F,0x2E); TEAL=RGBColor(0x2A,0x8A,0x8C)
WHITE=RGBColor(0xFF,0xFF,0xFF); INK=RGBColor(0x1F,0x2A,0x44); GRAY=RGBColor(0x6B,0x72,0x80); LIGHT=RGBColor(0xF2,0xF4,0xF8); MID=RGBColor(0xC9,0xD1,0xE0)
PAPER=RGBColor(0xFF,0xFF,0xFF)
W,H=13.333,7.5
HEAD="Georgia"; BODY="Calibri"
FOOT="EstimatingCoE | Eversource Energy | Stewards of Affordability"
def prs():
    p=Presentation(); p.slide_width=Inches(W); p.slide_height=Inches(H); return p
def blank(p): return p.slides.add_slide(p.slide_layouts[6])
def rect(s,x,y,w,h,fill,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(0.75)
    sh.shadow.inherit=False; return sh
def text(s,x,y,w,h,txt,size=18,bold=False,color=INK,font=BODY,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,italic=False,margin=0.05,line_spacing=1.05):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(0.03); tf.vertical_anchor=anchor
    lines=txt if isinstance(txt,list) else [txt]
    for i,ln in enumerate(lines):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.alignment=align; para.line_spacing=line_spacing
        runs=ln if isinstance(ln,list) else [ln]
        for rn in runs:
            if isinstance(rn,tuple): t,opt=rn
            else: t,opt=rn,{}
            r=para.add_run(); r.text=t; f=r.font; f.name=opt.get("font",font); f.size=Pt(opt.get("size",size)); f.bold=opt.get("bold",bold); f.italic=opt.get("italic",italic); f.color.rgb=opt.get("color",color)
        para.space_after=Pt(opt.get("after",4) if isinstance(rn,tuple) else 4)
    return tb
def bullets(s,x,y,w,h,items,size=18,color=INK,gap=6,bullet_color=GOLD):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.05)
    for i,it in enumerate(items):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph(); para.space_after=Pt(gap); para.line_spacing=1.05
        r=para.add_run(); r.text="▪  "; r.font.size=Pt(size); r.font.color.rgb=bullet_color; r.font.name=BODY
        if isinstance(it,(tuple,list)):
            r1=para.add_run(); r1.text=it[0]; r1.font.bold=True; r1.font.size=Pt(size); r1.font.color.rgb=color; r1.font.name=BODY
            r2=para.add_run(); r2.text=it[1]; r2.font.size=Pt(size); r2.font.color.rgb=color; r2.font.name=BODY
        else:
            r1=para.add_run(); r1.text=it; r1.font.size=Pt(size); r1.font.color.rgb=color; r1.font.name=BODY
    return tb
def chrome(s,title,eyebrow=None,n=None,dark=False):
    bg=NAVY if dark else PAPER
    rect(s,0,0,W,H,bg)
    tc=WHITE if dark else NAVY
    if eyebrow: text(s,0.6,0.28,11,0.3,eyebrow.upper(),size=11,bold=True,color=(GOLD if dark else BURG),font=BODY)
    text(s,0.6,0.52,12.1,0.95,title,size=23,bold=True,color=tc,font=HEAD,anchor=MSO_ANCHOR.TOP,line_spacing=1.0)
    text(s,0.6,7.02,9,0.3,FOOT,size=10,color=(MID if dark else GRAY))
    if n: text(s,12.0,7.02,0.75,0.3,str(n),size=10,color=(MID if dark else GRAY),align=PP_ALIGN.RIGHT)
def stat(s,x,y,w,h,value,label,sub=None,fill=LIGHT,vcolor=NAVY,vsize=40):
    rect(s,x,y,w,h,fill)
    text(s,x+0.15,y+0.12,w-0.3,0.9,value,size=vsize,bold=True,color=vcolor,font=HEAD)
    text(s,x+0.15,y+0.12+vsize/72*1.35,w-0.3,0.5,label,size=15,bold=True,color=INK)
    if sub: text(s,x+0.15,y+0.12+vsize/72*1.35+0.42,w-0.3,h-(0.12+vsize/72*1.35+0.42)-0.05,sub,size=12.5,color=GRAY,line_spacing=1.0)
def table(s,x,y,w,h,headers,rows,col_w=None,size=13,header_fill=NAVY,zebra=True,bold_last=False,align_right=()):
    nr,nc=len(rows)+1,len(headers)
    gt=s.shapes.add_table(nr,nc,Inches(x),Inches(y),Inches(w),Inches(h)); t=gt.table
    if col_w:
        for i,cw in enumerate(col_w): t.columns[i].width=Inches(cw)
    for j,hd in enumerate(headers):
        c=t.cell(0,j); c.fill.solid(); c.fill.fore_color.rgb=header_fill; c.text=""; p=c.text_frame.paragraphs[0]; r=p.add_run(); r.text=hd; r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name=BODY
        c.margin_left=c.margin_right=Inches(0.06); c.margin_top=c.margin_bottom=Inches(0.03)
        if j in align_right: p.alignment=PP_ALIGN.RIGHT
    for i,row in enumerate(rows,1):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.fill.solid(); c.fill.fore_color.rgb=(LIGHT if (zebra and i%2==0) else WHITE); c.text=""
            p=c.text_frame.paragraphs[0]; r=p.add_run(); r.text=str(v); r.font.size=Pt(size); r.font.color.rgb=INK; r.font.name=BODY
            r.font.bold = (bold_last and i==nr-1)
            c.margin_left=c.margin_right=Inches(0.06); c.margin_top=c.margin_bottom=Inches(0.03)
            if j in align_right: p.alignment=PP_ALIGN.RIGHT
            c.vertical_anchor=MSO_ANCHOR.MIDDLE
    return t
def notes(s,txt): s.notes_slide.notes_text_frame.text=txt
def style_chart(ch,legend=True,size=12,legend_pos=XL_LEGEND_POSITION.BOTTOM):
    ch.has_title=False
    ch.has_legend=legend
    if legend: ch.legend.position=legend_pos; ch.legend.include_in_layout=False; ch.legend.font.size=Pt(size); ch.legend.font.name=BODY
    ch.font.name=BODY; ch.font.size=Pt(size); ch.font.color.rgb=INK
